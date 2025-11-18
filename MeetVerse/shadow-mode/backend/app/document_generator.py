"""
Document generation (Word, PDF, PPT)
"""
from docx import Document as DocxDocument
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import os

OUTPUT_DIR = os.getenv("DOCUMENT_OUTPUT_DIR", "./documents")

os.makedirs(OUTPUT_DIR, exist_ok=True)

class DocumentGenerator:
    def generate(self, content: str, filename: str, doc_type: str = "docx"):
        """Generate document of specified type"""
        if doc_type == "docx":
            return self._generate_docx(content, filename)
        elif doc_type == "pdf":
            return self._generate_pdf(content, filename)
        elif doc_type == "pptx":
            return self._generate_pptx(content, filename)
        else:
            raise ValueError(f"Unsupported document type: {doc_type}")
    
    def _generate_docx(self, content: str, filename: str):
        """Generate Word document"""
        doc = DocxDocument()
        
        # Split content into paragraphs
        paragraphs = content.split("\n\n")
        for para in paragraphs:
            if para.strip():
                doc.add_paragraph(para.strip())
        
        file_path = os.path.join(OUTPUT_DIR, f"{filename}.docx")
        doc.save(file_path)
        return file_path
    
    def _generate_pdf(self, content: str, filename: str):
        """Generate PDF document"""
        file_path = os.path.join(OUTPUT_DIR, f"{filename}.pdf")
        doc = SimpleDocTemplate(file_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        
        paragraphs = content.split("\n\n")
        for para in paragraphs:
            if para.strip():
                story.append(Paragraph(para.strip(), styles["Normal"]))
                story.append(Spacer(1, 12))
        
        doc.build(story)
        return file_path
    
    def _generate_pptx(self, content: str, filename: str):
        """Generate PowerPoint presentation"""
        prs = Presentation()
        
        # Split content into slides (by double newlines or sections)
        sections = content.split("\n\n")
        for i, section in enumerate(sections):
            if section.strip():
                slide = prs.slides.add_slide(prs.slide_layouts[0])
                title = slide.shapes.title
                content_placeholder = slide.placeholders[1]
                
                title.text = f"Slide {i + 1}"
                content_placeholder.text = section.strip()
        
        file_path = os.path.join(OUTPUT_DIR, f"{filename}.pptx")
        prs.save(file_path)
        return file_path

